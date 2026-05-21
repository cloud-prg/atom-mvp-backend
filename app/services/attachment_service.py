from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile

from fastapi import UploadFile


MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 12000


@dataclass
class ParsedAttachment:
    filename: str
    content_type: str
    kind: str
    text: str


async def parse_uploads(files: list[UploadFile] | None) -> list[ParsedAttachment]:
    parsed: list[ParsedAttachment] = []
    for file in files or []:
        content = await file.read()
        parsed.append(parse_attachment(file.filename or "attachment", file.content_type or "", content))
    return parsed


def parse_attachment(filename: str, content_type: str, content: bytes) -> ParsedAttachment:
    if len(content) > MAX_ATTACHMENT_BYTES:
        return ParsedAttachment(filename, content_type, "unsupported", "附件超过 10MB，后端未解析。")

    suffix = Path(filename).suffix.lower()
    if _is_plain_text(suffix, content_type):
        return ParsedAttachment(filename, content_type, "text", _decode_text(content))
    if suffix == ".pdf" or content_type == "application/pdf":
        return ParsedAttachment(filename, content_type, "pdf", _extract_pdf_text(content))
    if suffix == ".docx" or content_type.endswith("wordprocessingml.document"):
        return ParsedAttachment(filename, content_type, "document", _extract_docx_text(content))
    if suffix in {".xlsx", ".xls"} or "spreadsheet" in content_type:
        return ParsedAttachment(filename, content_type, "spreadsheet", _extract_spreadsheet_text(content))
    if content_type.startswith("image/") or suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"}:
        return ParsedAttachment(filename, content_type, "image", "图片附件已收到；当前后端暂未启用视觉模型，只能把文件名和类型作为上下文。")

    return ParsedAttachment(filename, content_type, "unsupported", "当前文件类型暂不支持解析。")


def build_attachment_context(attachments: list[ParsedAttachment]) -> str:
    if not attachments:
        return ""
    sections = []
    for index, attachment in enumerate(attachments, start=1):
        text = attachment.text.strip() or "未提取到可读文本。"
        sections.append(
            f"[Attachment {index}: {attachment.filename} | {attachment.kind} | {attachment.content_type}]\n{text[:MAX_EXTRACTED_CHARS]}"
        )
    return "\n\n".join(sections)


def _is_plain_text(suffix: str, content_type: str) -> bool:
    return content_type.startswith("text/") or suffix in {".txt", ".md", ".csv", ".json", ".log"}


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return content.decode(encoding)[:MAX_EXTRACTED_CHARS]
        except UnicodeDecodeError:
            continue
    return ""


def _extract_pdf_text(content: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        return "PDF 解析依赖未安装，请安装 pypdf 后重试。"

    reader = PdfReader(BytesIO(content))
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    return text[:MAX_EXTRACTED_CHARS]


def _extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError:
        return "Word 解析依赖未安装，请安装 python-docx 后重试。"

    document = Document(BytesIO(content))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    return text[:MAX_EXTRACTED_CHARS]


def _extract_spreadsheet_text(content: bytes) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        return "Excel 解析依赖未安装，请安装 openpyxl 后重试。"

    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except BadZipFile:
        return "无法解析该表格文件，请确认文件格式是否正确。"

    lines: list[str] = []
    for sheet in workbook.worksheets[:5]:
        lines.append(f"# Sheet: {sheet.title}")
        for row in sheet.iter_rows(max_row=50, values_only=True):
            values = ["" if value is None else str(value) for value in row]
            if any(values):
                lines.append(",".join(values))
    return "\n".join(lines)[:MAX_EXTRACTED_CHARS]
